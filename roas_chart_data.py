import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ROAS Data"

coral = "E8541C"
navy  = "0B1929"
dark  = "0F2236"
white = "FFFFFF"
grey  = "7A8FA8"

hdr_fill   = PatternFill("solid", fgColor=coral)
row_fill1  = PatternFill("solid", fgColor="0F2236")
row_fill2  = PatternFill("solid", fgColor="0A1828")
thin = Border(
    left=Side(style='thin', color="1E3A50"),
    right=Side(style='thin', color="1E3A50"),
    top=Side(style='thin', color="1E3A50"),
    bottom=Side(style='thin', color="1E3A50"),
)

headers = ["Month", "Spend ($)", "Revenue Returned ($)", "ROAS"]
data    = [
    ["Jul 26", 1595, 16013, 10.04],
    ["Aug 26", 1219, 42574, 34.93],
]

ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 24
ws.column_dimensions["D"].width = 12

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font      = Font(name="Arial", bold=True, color=white, size=11)
    cell.fill      = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin

for r, row in enumerate(data, 2):
    fill = row_fill1 if r % 2 == 0 else row_fill2
    for col, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font      = Font(name="Arial", color=white, size=11)
        cell.fill      = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = thin
        if col == 2:
            cell.number_format = '$#,##0'
        elif col == 3:
            cell.number_format = '$#,##0'
        elif col == 4:
            cell.number_format = '0.00"×"'

ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 20

# Grouped bar chart — Spend vs Revenue
chart = BarChart()
chart.type        = "col"
chart.grouping    = "clustered"
chart.title       = "$2,814 In. $58,587 Back."
chart.y_axis.title = "Amount ($)"
chart.x_axis.title = ""
chart.shape       = 4
chart.width       = 18
chart.height      = 12

cats = Reference(ws, min_col=1, min_row=2, max_row=3)

spend_ref = Reference(ws, min_col=2, min_row=1, max_row=3)
spend_ser = Series(spend_ref, title_from_data=True)
spend_ser.graphicalProperties.solidFill   = "1E3A55"
spend_ser.graphicalProperties.line.solidFill = "1E3A55"
chart.append(spend_ser)

rev_ref = Reference(ws, min_col=3, min_row=1, max_row=3)
rev_ser = Series(rev_ref, title_from_data=True)
rev_ser.graphicalProperties.solidFill   = coral
rev_ser.graphicalProperties.line.solidFill = coral
chart.append(rev_ser)

chart.set_categories(cats)
chart.dataLabels = DataLabelList()
chart.dataLabels.showVal = True
chart.dataLabels.showLegendKey = False
chart.dataLabels.showCatName  = False
chart.dataLabels.showSerName  = False

ws.add_chart(chart, "A6")

wb.save("/home/user/dashbaord/GMS_Oakberry_ROAS_Chart.xlsx")
print("done")
